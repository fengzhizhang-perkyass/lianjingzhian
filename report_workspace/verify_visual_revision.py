from pathlib import Path
import json,re,hashlib,sys
from docx import Document
from pypdf import PdfReader
from PIL import Image,ImageDraw
R=Path(__file__).parent
Q=Path(sys.argv[1]) if len(sys.argv)>1 else R/'qa/visual_revision'
target=R.parent/'作品报告成品/Web3轻量化反洗钱智能分析系统_作品报告_图表优化版.docx'
doc=Document(target)
man=(R/'manuscript.md').read_text(encoding='utf-8').splitlines()
texts=[p.text for p in doc.paragraphs]
missing=[];headers=[];prose=[];figures=[]
for ln in man:
 ln=ln.strip()
 if not ln or ln.startswith('|') or ln.startswith('$$'):continue
 if ln.startswith('#'):
  s=ln.lstrip('#').strip();headers.append(s)
 elif ln.startswith('!['):
  s=re.match(r'!\[(.*?)\]',ln)[1];figures.append(s)
 else:s=ln;prose.append(s)
 if s not in texts:missing.append(s)
all_cells=[c.text for t in doc.tables for row in t.rows for c in row.cells]
missing_cells=[]
for ln in man:
 if ln.startswith('|') and not re.match(r'^\|[\s:|\-]+\|$',ln):
  for cell in [x.strip() for x in ln.strip('|').split('|')]:
   if cell not in all_cells:missing_cells.append(cell)
pdf=PdfReader(Q/'report.pdf')
pages=[]
for i,p in enumerate(pdf.pages,1):
 t=p.extract_text();pages.append({'page':i,'chars':len(t),'figures':re.findall(r'图\s*\d+\s*[-－]\s*\d+[^\n]*',t)})
out={'pages':len(pdf.pages),'figures':len(figures),'tables':len(doc.tables),'headings':len(headers),'prose_and_references':len(prose),'missing_paragraphs':missing,'missing_cells':missing_cells,'sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'page_index':pages}
(Q/'integrity.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in out.items() if k!='page_index'},ensure_ascii=False,indent=2))
ims=sorted(Q.glob('page-*.png'))
for k in range(0,len(ims),4):
 opened=[Image.open(p).convert('RGB') for p in ims[k:k+4]]
 w=max(i.width for i in opened);h=max(i.height for i in opened)
 sheet=Image.new('RGB',(2*w+30,2*(h+32)+25),'#D1D1D1');d=ImageDraw.Draw(sheet)
 for j,(p,im) in enumerate(zip(ims[k:k+4],opened)):
  x=j%2*(w+30);y=j//2*(h+57);d.text((x+8,y+5),p.stem,fill='black');sheet.paste(im,(x,y+30))
 sheet.save(Q/f'sheet-{k//4:02}.jpg',quality=92)

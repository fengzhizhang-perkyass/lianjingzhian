from pathlib import Path
from copy import deepcopy
import re,json
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

R=Path(__file__).parent
OUT=R.parent/'作品报告成品';OUT.mkdir(exist_ok=True)
TARGET=OUT/'Web3轻量化反洗钱智能分析系统_作品报告.docx'
d=Document(R/'template/template.docx')
# Preserve the actual template cover and mandatory instructions page.
paras=d.paragraphs
cut=paras[27]._p
started=False
for el in list(d._element.body):
 if el is cut:started=True
 if started and el.tag!=qn('w:sectPr'):d._element.body.remove(el)
for sect in d.sections:
 sect.page_width=Cm(21);sect.page_height=Cm(29.7)
 sect.top_margin=Cm(2.54);sect.bottom_margin=Cm(2.54)
 sect.left_margin=Cm(2.8575);sect.right_margin=Cm(2.5859)
 for tag in ['headerReference','footerReference']:
  for el in list(sect._sectPr.findall(qn('w:'+tag))):sect._sectPr.remove(el)
 for el in list(sect._sectPr.findall(qn('w:docGrid'))):sect._sectPr.remove(el)

def font(run,name='宋体',size=12,bold=False):
 run.font.name=name;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
 rf=run._element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),name)
 return run
for st in ['Normal','Body Text','Caption','TOC 1','TOC 2','Title','Heading 1','Heading 2','Heading 3']:
 try:s=d.styles[st]
 except KeyError:
  from docx.enum.style import WD_STYLE_TYPE
  s=d.styles.add_style(st,WD_STYLE_TYPE.PARAGRAPH)
 s.font.name='宋体';s.font.size=Pt(12);s.font.color.rgb=RGBColor(0,0,0)
 rf=s.element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),'宋体')
 pf=s.paragraph_format;pf.line_spacing=1.5;pf.space_after=Pt(0);pf.space_before=Pt(0)
 pf.widow_control=True
 if st.startswith('Heading'):
  n=int(st[-1]);s.font.name='黑体';rf.set(qn('w:eastAsia'),'黑体');s.font.size=Pt({1:16,2:14,3:12}[n]);s.font.bold=True
  pf.space_before=Pt(12 if n>1 else 0);pf.space_after=Pt(6);pf.keep_with_next=True
  pf.page_break_before=n==1;pf.first_line_indent=Pt(0)
  outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),str(n-1));s.element.get_or_add_pPr().append(outline)
  # Explicitly suppress inherited paragraph borders.
  for el in list(s.element.get_or_add_pPr().findall(qn('w:pBdr'))):s.element.get_or_add_pPr().remove(el)
d.styles['Normal'].paragraph_format.first_line_indent=Pt(24)
d.styles['TOC 1'].paragraph_format.first_line_indent=Pt(0)
d.styles['TOC 2'].paragraph_format.first_line_indent=Pt(0)
d.styles['TOC 2'].paragraph_format.left_indent=Pt(18)
d.styles['Title'].font.name='黑体';d.styles['Title'].font.size=Pt(26)

def replace(p,txt,size=16):
 p.clear();font(p.add_run(txt),'黑体',size,True)
paras=d.paragraphs
replace(paras[12],'作品名称：Web3轻量化反洗钱智能分析系统',16)
replace(paras[14],'提交日期：2026年10月13日',16)
for i in [3,4]:paras[i].style=d.styles['Title'];paras[i].paragraph_format.first_line_indent=Pt(0)
for i in range(16,22):
 paras[i].paragraph_format.line_spacing=1.5
 paras[i].paragraph_format.first_line_indent=Pt(0)
 if i>16:
  for run in paras[i].runs:font(run)

# TOC has its own section inherited from the template.
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
font(p.add_run('目  录'),'黑体',16,True)
p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');p._p.append(fld)
d.add_section(WD_SECTION_START.NEW_PAGE)
bodysect=d.sections[-1]
bodysect.header.is_linked_to_previous=False;bodysect.footer.is_linked_to_previous=False
hp=bodysect.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.CENTER;hp.paragraph_format.first_line_indent=Pt(0)
font(hp.add_run('Web3轻量化反洗钱智能分析系统  作品报告'),size=12)
fp=bodysect.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.paragraph_format.first_line_indent=Pt(0)
font(fp.add_run('— '))
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
font(fp.add_run(' —'))
pn=OxmlElement('w:pgNumType');pn.set(qn('w:start'),'1');bodysect._sectPr.append(pn)

def paragraph(txt,style=None):
 p=d.add_paragraph(style=style)
 font(p.add_run(txt))
 p.paragraph_format.line_spacing=1.5
 if style is None:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 return p
def caption(txt):
 p=paragraph(txt,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
 p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(6)
 return p
table_count={};chapter=0;image_count=0;table_total=0;head=''
table_titles=[
'背景数据及其统计口径','三个模块的输出契约','当前联合模型训练配置','追踪默认约束参数',
'分层测试对象与验证内容','测试环境','数据处理规模','当前模型分类与排序结果','已知测试地址的预算命中',
'严格移除实验结果','构造图目标恢复','相对原Beam的真实图结果','覆盖与在线开销对比','长链模式探索性复测',
'报告协议故障与修复测试','两次真实DeepSeek报告验收','端到端验收记录','痛点 创新与实验对应','技术机制对照','后续验证与试用安排']
def table(rows):
 global table_total
 table_total+=1;table_count[chapter]=table_count.get(chapter,0)+1
 cp=caption(f'表{chapter}-{table_count[chapter]} '+table_titles[table_total-1]);cp.paragraph_format.keep_with_next=True
 cols=len(rows[0]);t=d.add_table(rows=1,cols=cols);t.autofit=False
 widths={2:[7.3,8.25],3:[4.35,4.0,7.2],4:[6.1,3.15,3.15,3.15],5:[5.0,2.6,2.6,2.65,2.7]}[cols]
 # tailor numeric tables for short metrics, keep long method labels readable.
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 props=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for k in ['top','left','bottom','right','insideH','insideV']:
  b=OxmlElement('w:'+k);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
 props.append(borders)
 for i,row in enumerate(rows):
  rr=t.rows[0] if i==0 else t.add_row()
  trp=rr._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trp.append(cant)
  if i==0:rep=OxmlElement('w:tblHeader');trp.append(rep)
  for j,txt in enumerate(row):
   c=rr.cells[j];c.width=Cm(widths[j]);tc=c._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DBE7F2' if i==0 else ('F4F7FA' if i%2==0 else 'FFFFFF'));tc.append(sh)
   mar=OxmlElement('w:tcMar')
   for key,val in [('top',65),('bottom',65),('left',95),('right',95)]:
    z=OxmlElement('w:'+key);z.set(qn('w:w'),str(val));z.set(qn('w:type'),'dxa');mar.append(z)
   tc.append(mar)
   va=OxmlElement('w:vAlign');va.set(qn('w:val'),'center');tc.append(va)
   p=c.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1.5
   p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=(i==0)
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or (j>0 and len(txt)<18) else WD_ALIGN_PARAGRAPH.LEFT
   font(p.add_run(txt),bold=i==0)
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1;p.paragraph_format.first_line_indent=Pt(0);font(p.add_run(''),size=12)

equations=[
'z(v,d,j) = clip[(x(v,j) − μ(d,j)) / max(σ(d,j), 0.5), −8, 8]',
's(v) = 100 × sigmoid[(m(v) − log(w₊)) / T]',
'm(v) = b + Φ行为(v) + Φ偏离(v) + Φ结构(v)']
eqidx=0
lines=(R/'manuscript.md').read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line:continue
 if line.startswith('#'):
  level=len(line)-len(line.lstrip('#'));title=line[level:].strip();head=title
  if level==1:
   for num,ch in enumerate(['一','二','三','四','五'],1):
    if title.startswith('第'+ch+'章'):chapter=num
  p=d.add_paragraph(title,style=f'Heading {level}')
  ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),str(level-1));p._p.get_or_add_pPr().append(ol)
  if title=='摘要':p.paragraph_format.page_break_before=False
  for run in p.runs:font(run,'黑体',{1:16,2:14,3:12}[level],True)
  continue
 if line.startswith('!['):
  m=re.match(r'!\[(.*?)\]\((.*?)\)',line);title,path=m.groups();path=(R/path).resolve()
  if not path.exists():raise FileNotFoundError(path)
  im=Image.open(path);w,h=im.size;wc=15.55;hc=wc*h/w
  if hc>13:wc=13*w/h
  p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
  p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5);p.paragraph_format.line_spacing=1
  p.add_run().add_picture(str(path),width=Cm(wc))
  caption(title);image_count+=1;continue
 if line.startswith('|'):
  rr=[line]
  while i<len(lines) and lines[i].strip().startswith('|'):rr.append(lines[i].strip());i+=1
  rows=[[x.strip() for x in l.strip('|').split('|')] for l in rr if not re.match(r'^\|[\s:|\-]+\|$',l)]
  table(rows);continue
 if line.startswith('$$'):
  p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
  math=OxmlElement('m:oMath');run=OxmlElement('m:r');tx=OxmlElement('m:t');tx.text=equations[eqidx];eqidx+=1;run.append(tx);math.append(run);p._p.append(math)
  continue
 p=paragraph(line)
 if line.startswith('['):
  p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(6)
  p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.keep_together=True
 if line.startswith('关键词'):p.paragraph_format.first_line_indent=Pt(0)

# Do not leak local account names or revision history in core properties.
d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title='Web3轻量化反洗钱智能分析系统作品报告'
d.core_properties.subject='中国研究生网络安全创新大赛作品报告';d.core_properties.comments=''
settings=d.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
d.save(TARGET)
(R/'build_stats.json').write_text(json.dumps({'output':str(TARGET),'figures':image_count,'tables':table_total,'characters':len(''.join(lines))},ensure_ascii=False,indent=2),encoding='utf-8')
print(TARGET);print('figures',image_count,'tables',table_total)
